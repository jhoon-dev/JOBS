"""Read exports from jobs_report_tests independently; never rewrite them.

Optional development/CI dependencies: openpyxl==3.1.5, pypdf==6.10.0.
Usage: python scripts/verify-report-fixtures.py build
"""
import re
import sys
import zipfile
import unicodedata
from pathlib import Path

from openpyxl import load_workbook
from pypdf import PdfReader


def verify(root: Path):
    for name in ("reports-smoke.xlsx", "reports-resources.xlsx", "reports-pdm.xlsx", "reports-constraints.xlsx", "reports-effort.xlsx"):
        with zipfile.ZipFile(root / name) as archive:
            assert archive.testzip() is None, "ZIP checksum failure"
            assert not any("externalLink" in p or "vbaProject" in p for p in archive.namelist())
        book = load_workbook(root / name, data_only=False)
        assert "출처와 설정" in book.sheetnames
        for sheet in book:
            assert sheet.freeze_panes == "A7"
            assert sheet.auto_filter.ref.startswith("A6:")
            assert not any(cell.data_type == "f" for row in sheet for cell in row)
        if name == "reports-smoke.xlsx":
            assert len(book.sheetnames) == 14
            costs = book["작업별 원가"]
            assert costs["G7"].value == 6 and costs["G7"].data_type == "n"
            assert costs["G8"].value == 0 and costs["G8"].data_type == "n"
            task = book["계획 일정"]["B7"]
            assert task.value.startswith("=HYPERLINK") and task.data_type == "s"
            assert book["실적과 예측"]["C7"].value == "미정"
            assert book["기간별 원가"]["G7"].value == 6
            assert book["기간별 원가"]["G8"].value == "미정"
            metadata = dict(book["출처와 설정"].iter_rows(min_row=7, values_only=True))
            assert metadata["보고서 ID"] == "report-contract-fixture"
            assert metadata["원본 상태"] == "미저장 메모리 revision"
        elif name == "reports-effort.xlsx":
            assert len(book.sheetnames) == 19
            work = book["공수와 작업 유형"]
            assert work["C7"].value == "고정 공수" and work["D7"].value == "공수"
            assert work["E7"].value == work["F7"].value == 8
            assert work["E7"].data_type == work["F7"].data_type == "n"
            assert work["G7"].value == 13
            assert book["계획 일정"]["F7"].value == "2026-10-28 15:00"
            assert book["계획 일정"]["D7"].value == 780
            assert book["분할 작업 중단"]["C7"].value == "2026-10-26 13:00"
            assert book["자원별 근무 달력"]["D7"].value == "Designer Tuesday off"
            spans = list(book["실행 배정 구간"].iter_rows(min_row=7, values_only=True))
            assert len(spans) == 3 and sum(row[5] for row in spans) == 8
            assert not any("2026-10-27" in str(row[2]) for row in spans)
            assert book["작업별 원가"]["C7"].value == 48000
        elif name == "reports-constraints.xlsx":
            assert len(book.sheetnames) == 15
            timing = book["작업 제약과 마감"]
            assert timing["C7"].value == "SNET · 시작 이후"
            assert timing["D7"].value == "2026-10-27 09:00" and timing["E7"].value == "2026-10-27 17:00"
            assert timing["G7"].value == -9 and timing["G7"].data_type == "n"
            assert timing["H7"].value == "마감 초과"
        elif name == "reports-pdm.xlsx":
            assert len(book.sheetnames) == 14
            links = book["실행 선행관계"]
            assert links["E7"].value == "SS" and links["F7"].value == -60
            assert links["F7"].data_type == "n" and links["G7"].value == "PDM Afternoon"
            assert book["계획 일정"]["I8"].value == "PDM Afternoon"
        else:
            assert "작업별 원가" not in book.sheetnames and "자원 등록과 가용량" in book.sheetnames
        print(f"PASS: {name}: {len(book.sheetnames)} sheets, ZIP/OOXML/types/unknown/provenance")

    for name in ("reports-smoke.pdf", "reports-long-portrait.pdf", "reports-a3.pdf", "reports-pdm.pdf", "reports-constraints.pdf", "reports-effort.pdf"):
        reader = PdfReader(root / name, strict=True)
        raw = "\n".join(p.extract_text() or "" for p in reader.pages)
        assert not any(unicodedata.category(c) == "Cc" and c not in "\n\r\t" for c in raw), f"PDF font emitted control characters: {name}"
        # Native PDF fonts can encode typographic ligatures (e.g. ff as U+FB00).
        # Normalize presentation forms before comparing semantic text/numbers.
        compact = re.sub(r"\s+", "", unicodedata.normalize("NFKC", raw))
        assert reader.pages and all("JOBS" in (p.extract_text() or "") for p in reader.pages)
        if name == "reports-long-portrait.pdf":
            numbers = [int(n) for n in re.findall(r"Segment_(\d+)", compact)]
            assert len(numbers) == 700 and set(numbers) == set(range(700)), "long row lost/duplicated content"
            assert "END_SENTINEL_739" in compact and len(reader.pages) >= 3
            assert abs(float(reader.pages[0].mediabox.width) - 595) < 3
            assert abs(float(reader.pages[0].mediabox.height) - 842) < 3
        elif name == "reports-effort.pdf":
            assert "effort-report-contract" in compact and "고정공수" in compact
            assert "DesignerTuesdayoff" in compact, f"resource calendar text missing: {compact!r}"
            assert "2026-10-2815:00" in compact, f"effort finish stamp missing: {compact!r}"
            assert "실행배정구간" in compact and "48000" in compact
        elif name == "reports-constraints.pdf":
            assert "constraints-report-contract" in compact and "SNET" in compact
            assert "-9.00" in compact and "마감초과" in compact
        elif name == "reports-pdm.pdf":
            assert "pdm-report-contract" in compact and "PDMAfternoon" in compact
            assert "SS" in compact and "-60" in compact
        else:
            assert "report-contract-fixture" in compact and "6.00" in compact
            width = 1191 if name == "reports-a3.pdf" else 842
            assert abs(float(reader.pages[0].mediabox.width) - width) < 3
        print(f"PASS: {name}: {len(reader.pages)} pages, size/text/provenance/continuation")


if __name__ == "__main__":
    verify(Path(sys.argv[1] if len(sys.argv) > 1 else "build"))
